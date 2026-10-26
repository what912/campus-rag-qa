"""Build the editable technical draft using the bundled document runtime."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]


def set_font(style,name,size):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)


def add_table(document,lines):
    rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-:]+',v) for v in r)]
    table=document.add_table(rows=0,cols=len(rows[0]));table.autofit=False
    widths=[3.8]+[(6.5-3.8)/(len(rows[0])-1)]*(len(rows[0])-1) if len(rows[0])==3 else [3.3,3.2]
    for i,values in enumerate(rows):
        cells=table.add_row().cells
        for j,value in enumerate(values):
            cell=cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1.15
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(value);r.font.size=Pt(10)
            if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
            properties=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173F52' if i==0 else ('F1F5F7' if i%2==0 else 'FFFFFF'));properties.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            properties.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'95');m.set(qn('w:type'),'dxa');margins.append(m)
            properties.append(margins)
        if i==0:
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    document.add_paragraph()


if __name__=='__main__':
    document=Document();section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.85)
    section.left_margin=section.right_margin=Inches(1)
    set_font(document.styles['Normal'],'宋体',11)
    normal=document.styles['Normal'].paragraph_format;normal.line_spacing=1.25;normal.space_after=Pt(6)
    for name,size in [('Title',19),('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
        set_font(document.styles[name],'微软雅黑',size)
        document.styles[name].paragraph_format.keep_with_next=True
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('技术草稿  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for r in footer.runs:r.font.size=Pt(9)
    lines=(ROOT/'docs/论文技术初稿.md').read_text(encoding='utf-8').splitlines()
    table_lines=[]
    for line in lines:
        if line.startswith('|'):
            table_lines.append(line);continue
        if table_lines:
            add_table(document,table_lines);table_lines=[]
        if not line.strip():continue
        if line.startswith('# '):
            document.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):
            text=line[3:]
            p=document.add_paragraph(text,'Heading 1')
            if text.startswith('第'):p.paragraph_format.page_break_before=True
        elif line.startswith('### '):
            text=line[4:];document.add_paragraph(text,'Heading 2')
            image={'3.3 总体架构':'system-architecture.png','3.4 数据模型':'data-model.png','5.2 开发题与检索指标':'retrieval-optimization.png'}.get(text)
            if image:
                p=document.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(ROOT/'output/figures'/image),width=Inches(6.45))
                caption=document.add_paragraph({'system-architecture.png':'系统处理与问答流程','data-model.png':'知识库与评分数据模型','retrieval-optimization.png':'开发题检索检查  不代表生成准确率'}[image])
                caption.alignment=WD_ALIGN_PARAGRAPH.CENTER;caption.paragraph_format.space_after=Pt(12)
                for r in caption.runs:r.font.size=Pt(10)
        else:
            document.add_paragraph(line)
    if table_lines:add_table(document,table_lines)
    destination=ROOT/'output/论文技术初稿.docx';destination.parent.mkdir(exist_ok=True)
    document.save(destination)
    print(destination)
