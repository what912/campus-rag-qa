import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app import create_app
from campus.documents import chunk_pages, parse_document
from campus.providers import generate_answer, build_messages


class SystemTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.application = create_app(Path(self.temp.name), seed=False)
        self.client = TestClient(self.application)

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def upload(self, text='# 太空摄影课程\n\n## 考核\n太空摄影课程考核采用观测报告，占总成绩的 60%。', **fields):
        return self.client.post('/api/documents', files={'file': ('test.md', text.encode(), 'text/markdown')}, data=fields)

    def test_upload_retrieve_source_and_delete(self):
        result = self.upload(category='课程大纲', version='2026 演示')
        self.assertEqual(result.status_code, 201, result.text)
        document_id = result.json()['id']
        answer = self.client.post('/api/ask', json={'question': '太空摄影课程考核采用什么？'}).json()
        self.assertEqual(answer['mode'], 'extractive')
        self.assertTrue(any('60%' in source['text'] for source in answer['sources']))
        self.assertEqual(answer['sources'][0]['document_id'], document_id)
        self.assertIsNone(answer['sources'][0]['page'])
        self.assertTrue((Path(self.temp.name) / 'vectors.npz').exists())
        download = self.client.get('/api/documents/' + document_id + '/download')
        self.assertEqual(download.status_code, 200)
        self.assertIn('60%', download.content.decode())
        self.assertEqual(self.client.delete('/api/documents/' + document_id).status_code, 200)
        self.assertEqual(self.client.get('/api/documents/' + document_id).status_code, 404)
        self.assertEqual(self.client.post('/api/ask', json={'question': '太空摄影考核'}).json()['mode'], 'no_evidence')

    def test_duplicates_invalid_formats_and_blank_questions(self):
        self.assertEqual(self.upload().status_code, 201)
        self.assertEqual(self.upload().status_code, 400)
        response = self.client.post('/api/documents', files={'file': ('a.exe', b'123')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.post('/api/ask', json={'question': '   '}).status_code, 400)

    def test_filters_and_no_evidence(self):
        self.upload(category='课程大纲')
        result = self.client.post('/api/ask', json={'question': '太空摄影考核', 'category': '办事流程'}).json()
        self.assertEqual(result['mode'], 'no_evidence')
        result = self.client.post('/api/ask', json={'question': '红烧牛肉烹饪温度'}).json()
        self.assertEqual(result['sources'], [])

    def test_restart_preserves_documents_without_readding_samples(self):
        document_id = self.upload().json()['id']
        reopened = create_app(Path(self.temp.name), seed=False)
        with TestClient(reopened) as client:
            self.assertEqual(client.get('/api/status').json()['documents'], 1)
            self.assertEqual(client.get('/api/documents/' + document_id).status_code, 200)

    def test_settings_do_not_return_or_persist_credentials(self):
        response = self.client.put('/api/settings', json={
            'enabled': True, 'base_url': 'https://example.invalid/v1',
            'model': 'test-model', 'api_key': 'test-secret-only'})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('test-secret-only', self.client.get('/api/settings').text)
        self.assertNotIn('test-secret-only', (Path(self.temp.name) / 'settings.json').read_text(encoding='utf-8'))
        self.assertTrue(self.client.get('/api/settings').json()['has_key'])
        bad = self.client.put('/api/settings', json={'embedding_mode': 'semantic', 'model_path': 'missing'})
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.client.get('/api/status').json()['mode'], 'tfidf')

    def test_rag_mode_receives_evidence_and_failures_are_visible(self):
        self.upload(category='课程大纲')
        self.client.put('/api/settings', json={'enabled': True, 'base_url': 'https://example.invalid/v1', 'model': 'test'})
        with patch('app.generate_answer', return_value=('观测报告占 60%。[1]', [])) as generated:
            result = self.client.post('/api/ask', json={'question': '太空摄影课程考核'}).json()
            self.assertEqual(result['mode'], 'rag')
            self.assertTrue(any('60%' in source['text'] for source in generated.call_args.args[1]))
        with patch('app.generate_answer', side_effect=ValueError('服务不可用')):
            response = self.client.post('/api/ask', json={'question': '太空摄影课程考核'})
            self.assertEqual(response.status_code, 502)

    def test_example_labels_and_excluding_samples(self):
        seeded = create_app(Path(self.temp.name), seed=True)
        with TestClient(seeded) as client:
            result = client.post('/api/ask', json={'question': '在读证明怎么办理？'}).json()
            self.assertTrue(any('虚构' in warning for warning in result['warnings']))
            excluded = client.post('/api/ask', json={'question': '在读证明怎么办理？', 'include_samples': False}).json()
            self.assertEqual(excluded['mode'], 'no_evidence')
            evaluation = client.get('/api/evaluation?dataset=samples').json()
            self.assertEqual(len(evaluation['rows']), 10)
            for doc in client.get('/api/documents').json():
                client.delete('/api/documents/' + doc['id'])
        reopened = create_app(Path(self.temp.name), seed=True)
        with TestClient(reopened) as client:
            self.assertEqual(client.get('/api/status').json()['documents'], 0)


class ParsingTests(unittest.TestCase):
    def test_chunk_overlap_and_heading_page_boundaries(self):
        text = '# 一\n' + '甲乙丙丁。' * 300 + '\n# 二\n第二段资料。'
        chunks = chunk_pages([{'page': 1, 'text': text}, {'page': 2, 'text': '第三页示例文本。'}])
        self.assertTrue(all(len(c['text']) <= 400 for c in chunks))
        self.assertEqual({c['page'] for c in chunks}, {1, 2})
        self.assertFalse(any('# 二' in c['text'] and '# 一' in c['text'] for c in chunks))
        self.assertIn(chunks[0]['text'][-55:], chunks[1]['text'])

    def test_pdf_page_numbers_and_scanned_pdf_error(self):
        writer = PdfWriter()
        for text in ('Campus first page', 'Campus second page'):
            page = writer.add_blank_page(width=612, height=792)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                     NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            stream.set_data(('BT /F1 12 Tf 50 700 Td (' + text + ') Tj ET').encode())
            page[NameObject('/Contents')] = writer._add_object(stream)
        content = io.BytesIO()
        writer.write(content)
        pages = parse_document('test.pdf', content.getvalue())
        self.assertEqual([p['page'] for p in pages], [1, 2])
        self.assertIn('second page', pages[1]['text'])
        blank = PdfWriter()
        blank.add_blank_page(width=100, height=100)
        content = io.BytesIO()
        blank.write(content)
        with self.assertRaisesRegex(ValueError, 'OCR'):
            parse_document('scanned.pdf', content.getvalue())

    def test_docx_tables_and_gb18030_text(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('课程大纲')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '考核规则'
        doc.add_paragraph('末尾说明')
        content = io.BytesIO()
        doc.save(content)
        text = parse_document('example.docx', content.getvalue())[0]['text']
        self.assertLess(text.index('课程大纲'), text.index('考核规则'))
        self.assertLess(text.index('考核规则'), text.index('末尾说明'))
        self.assertEqual(parse_document('gb.txt', '学校资料'.encode('gb18030'))[0]['text'], '学校资料')


class ProviderTests(unittest.TestCase):
    def test_http_protocol_and_citation_validation(self):
        sources = [{'citation': 1, 'name': 'test.md', 'version': '', 'page': None,
                    'ordinal': 1, 'sample': True, 'text': '示例内容'}]
        self.assertIn('示例内容', build_messages('问题', sources)[1]['content'])
        transport = httpx.MockTransport(lambda request: httpx.Response(200, json={
            'choices': [{'message': {'content': '回答。[99]'}}]}))
        real_client = httpx.Client
        with patch('campus.providers.httpx.Client', side_effect=lambda **kwargs: real_client(transport=transport, **kwargs)):
            answer, warnings = generate_answer('问题', sources, {
                'base_url': 'https://example.invalid/v1', 'model': 'test', 'api_key': 'test'})
            self.assertIn('[99]', answer)
            self.assertTrue(any('不存在' in warning for warning in warnings))


if __name__ == '__main__':
    unittest.main()
